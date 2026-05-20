"""Excel parser using openpyxl — handles marketing results spreadsheets."""
from __future__ import annotations

import io
from typing import Any, Dict, List

try:
    import openpyxl
    OPENPYXL_OK = True
except ImportError:
    OPENPYXL_OK = False


def parse_excel(content: bytes) -> Dict[str, Any]:
    """
    Parse .xlsx bytes into rows grouped by sheet name.

    Returns:
        {
            "sheets": { sheet_name: { headers, rows, row_count } },
            "text": str,
        }
    """
    if not OPENPYXL_OK:
        return {"sheets": {}, "text": "[openpyxl not installed]"}

    wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    sheets: Dict[str, Any] = {}
    text_parts: List[str] = []

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        all_rows = list(ws.iter_rows(values_only=True))
        if not all_rows:
            continue
        headers = [str(c) if c is not None else "" for c in all_rows[0]]
        data_rows = []
        for row in all_rows[1:51]:     # cap at 50 data rows
            data_rows.append({headers[i]: row[i] for i in range(len(headers)) if i < len(row)})

        sheets[sheet_name] = {
            "headers": headers,
            "rows": data_rows,
            "row_count": len(all_rows) - 1,
        }
        text_parts.append(f"Sheet: {sheet_name}\n")
        text_parts.append(", ".join(headers) + "\n")
        for row in data_rows[:20]:
            text_parts.append(", ".join(str(v) for v in row.values()) + "\n")

    return {
        "sheets": sheets,
        "text": "".join(text_parts),
    }
