"""DOCX parser using python-docx — handles Word documents."""
from __future__ import annotations

import io
from typing import Any, Dict

try:
    from docx import Document
    DOCX_OK = True
except ImportError:
    DOCX_OK = False


def parse_docx(content: bytes) -> Dict[str, Any]:
    """
    Extract text and table data from .docx bytes.

    Returns:
        {
            "paragraphs": int,
            "text": str,
            "tables": List[List[List[str]]],
        }
    """
    if not DOCX_OK:
        return {"paragraphs": 0, "text": "[python-docx not installed]", "tables": []}

    doc = Document(io.BytesIO(content))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    tables = []
    for table in doc.tables:
        tbl_data = []
        for row in table.rows:
            tbl_data.append([cell.text.strip() for cell in row.cells])
        tables.append(tbl_data)

    return {
        "paragraphs": len(paragraphs),
        "text": "\n".join(paragraphs),
        "tables": tables[:10],
    }
